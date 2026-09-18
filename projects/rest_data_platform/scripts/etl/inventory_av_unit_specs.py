#!/usr/bin/env python3
"""Stufe-0 Inventar: 4C Specs + VPI Measurements + MaStR AV01-Payload."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ingest_mastr import build_catalog_lookup, iter_xml_records

REF = Path(
    r"C:\Users\he5013\academiccloudsync\ASTRA_Maritime_Circularity"
    r"\02_Data\Reference_Data"
)
MASTR_DIR = Path(
    r"C:\Users\he5013\academiccloudsync\ASTRA_Maritime_Circularity"
    r"\02_Data\Extraction\MaStR"
)
OUT = Path(
    r"C:\Users\he5013\ai-projects\projects\rest_data_platform"
    r"\reference\imc\av_unit_spec_inventory_2026_09.md"
)

TURBINE_XLSX = REF / "Offshore Wind Turbine Database.xlsx"
VPI_XLSX = REF / "Vessel & Ports Intelligence Database.xlsx"
AV01_SEE = "SEE982527987333"
WINDFARM_ID = "DE01"

COVERED_SPEC = {
    "Manufacturer",
    "TurbineName",
    "RatedPowerMW",
    "RotorDiameterm",
    "TowerHeightm",
    "TurbineId",
}


def cell(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def load_sheet(path: Path, sheet: str) -> tuple[list[str], list[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    out: list[dict] = []
    for row in rows:
        rec = {header[i]: row[i] if i < len(row) else None for i in range(len(header)) if header[i]}
        if any(v is not None and str(v).strip() for v in rec.values()):
            out.append(rec)
    wb.close()
    return header, out


def md_escape(value) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "/").replace("\n", " ")[:80]


def filled_table(title: str, rows: list[dict], keys: list[str], lines: list[str]) -> None:
    lines.append(f"### {title}")
    lines.append("")
    header = "| Feld |" + "".join(f" {r.get('_label', i)} |" for i, r in enumerate(rows))
    sep = "|------|" + "".join("---|" for _ in rows)
    lines.append(header)
    lines.append(sep)
    for key in keys:
        if key.startswith("_"):
            continue
        cells = []
        for rec in rows:
            val = rec.get(key)
            cells.append("—" if cell(val) is None else md_escape(val))
        lines.append(f"| `{key}` |" + "".join(f" {c} |" for c in cells))
    lines.append("")


def main() -> int:
    link_headers, links = load_sheet(TURBINE_XLSX, "Turbine on Windfarms")
    av_links = [r for r in links if str(r.get("WindfarmId") or "") == WINDFARM_ID]
    tids = [str(r.get("TurbineId")) for r in av_links if r.get("TurbineId") is not None]

    spec_headers, specs = load_sheet(TURBINE_XLSX, "Offshore Wind Turbine Specs")
    spec_by_id = {str(r.get("TurbineId")): r for r in specs if r.get("TurbineId") is not None}
    av_specs = []
    for tid in tids:
        rec = dict(spec_by_id.get(tid, {"TurbineId": tid}))
        rec["_label"] = f"tid {tid}"
        av_specs.append(rec)

    meas_headers, meas = load_sheet(VPI_XLSX, "Turbine Measurements")
    meas_by_id = {str(r.get("TurbineId")): r for r in meas if r.get("TurbineId") is not None}
    av_meas = []
    for tid in tids:
        rec = dict(meas_by_id.get(tid, {"TurbineId": tid}))
        rec["_label"] = f"tid {tid}"
        av_meas.append(rec)

    found_headers, founds = load_sheet(VPI_XLSX, "Fixed Foundation Measurements")
    av_founds = [r for r in founds if str(r.get("WindfarmId") or "") == WINDFARM_ID]

    catalog = build_catalog_lookup(MASTR_DIR / "Katalogwerte.xml")
    av01: dict | None = None
    av_units: list[dict] = []
    key_counts: dict[str, int] = {}
    for _, payload in iter_xml_records(MASTR_DIR / "EinheitenWind.xml", "EinheitWind"):
        name = str(payload.get("NameStromerzeugungseinheit") or "")
        park = str(payload.get("NameWindpark") or "")
        if name.startswith("AV") and "alpha ventus" in park.casefold():
            av_units.append(payload)
            for k in payload:
                key_counts[k] = key_counts.get(k, 0) + 1
            if payload.get("EinheitMastrNummer") == AV01_SEE:
                av01 = payload
        if len(av_units) >= 12 and av01 is not None:
            break

    lines: list[str] = []
    w = lines.append
    w("<!-- Reality Block")
    w("last_update: 2026-09-18")
    w("status: draft")
    w("scope:")
    w('  summary: "Stufe-0 Inventar Alpha Ventus: 4C Specs, VPI Measurements, MaStR AV01."')
    w("  in_scope:")
    w("    - filled vs empty fields for two AV turbine types")
    w("    - MaStR EinheitWind keys on AV01")
    w("  out_of_scope:")
    w("    - schema / ETL")
    w("    - inventing unit-to-type mapping")
    w("notes:")
    w('  - "Excel: academiccloudsync Reference_Data; MaStR EinheitenWind.xml UTF-16."')
    w("-->")
    w("")
    w("# AV Stufe-0 — Specs / Measurements / MaStR")
    w("")
    w("Quelle: 4C Turbine-DB + VPI + MaStR-XML (lokal). **Kein Schema.**")
    w("")
    w("## 4C Typ am Park (`Turbine on Windfarms`, WindfarmId DE01)")
    w("")
    w("| TurbineId | OEM | Modell | MW min/max | n min/max | Foundation | Expired |")
    w("|-----------|-----|--------|------------|-----------|------------|---------|")
    for rec in av_links:
        w(
            "| "
            + " | ".join(
                [
                    md_escape(rec.get("TurbineId")),
                    md_escape(rec.get("TurbineManufacturer")),
                    md_escape(rec.get("TurbineModel")),
                    f"{md_escape(rec.get('TurbineMWMin'))}/{md_escape(rec.get('TurbineMWMax'))}",
                    f"{md_escape(rec.get('NoTurbinesMin'))}/{md_escape(rec.get('NoTurbinesMax'))}",
                    md_escape(rec.get("Foundation") or rec.get("FoundationsOnWindfarm")),
                    md_escape(rec.get("IsExpired")),
                ]
            )
            + " |"
        )
    w("")
    w(
        f"Zeilen DE01: **{len(av_links)}**. 4C liefert **keine WEA-IDs** "
        "und **kein n=6+6** — nur zwei Typ-Zeilen am Park."
    )
    w("")

    spec_keys = [k for k in spec_headers if k]
    filled_table("4C Specs (`Offshore Wind Turbine Specs`)", av_specs, spec_keys, lines)

    already = COVERED_SPEC
    weight_keys = [k for k in spec_keys if "Weight" in k or k.startswith("Blades") or k.startswith("Nacelle") or k.startswith("Tower") or k.startswith("Gear") or k.startswith("Generator")]
    w("**Schon in `imc_turbine_models`:** Manufacturer, TurbineName, RatedPowerMW, RotorDiameterm, TowerHeightm (Nabe, oft leer).")
    w("")
    w("**Gewicht/BOM-relevant in Specs:**")
    w("")
    w("| Feld | tid " + " | tid ".join(tids) + " |")
    w("|------|" + "|".join("---" for _ in tids) + "|")
    for key in [
        "RotorWeightt",
        "BladesNumber",
        "BladesLengthm",
        "BladesMaterial",
        "BladesType",
        "BladesWeightt",
        "TowerType",
        "TowerHeightm",
        "TowerWeightt",
        "NacelleHeightm",
        "NacelleLengthm",
        "NacelleWidthm",
        "NacelleWeightt",
        "GearboxType",
        "GeneratorType",
    ]:
        cells = []
        for rec in av_specs:
            val = rec.get(key)
            cells.append("—" if cell(val) is None else md_escape(val))
        w(f"| `{key}` |" + "|".join(f" {c} " for c in cells) + "|")
    w("")

    meas_keys = [k for k in meas_headers if k]
    filled_table("VPI `Turbine Measurements`", av_meas, meas_keys, lines)

    w("## VPI Fixed Foundation (WindfarmId DE01)")
    w("")
    if not av_founds:
        w("Keine Zeile mit WindfarmId DE01.")
    else:
        w(f"{len(av_founds)} Zeile(n). Erste Keys/Werte (non-null):")
        w("")
        for rec in av_founds[:3]:
            w("| Feld | Wert |")
            w("|------|------|")
            for k, v in rec.items():
                if cell(v) is None:
                    continue
                w(f"| `{k}` | {md_escape(v)} |")
            w("")
    w("")

    w("## MaStR EinheitWind — AV01")
    w("")
    w(f"Gesucht: `{AV01_SEE}`. AV-Einheiten im Scan: **{len(av_units)}**.")
    w("")
    if av01 is None:
        w("AV01-Payload **nicht** gefunden.")
    else:
        w("Katalog-IDs wo auflösbar als `*_label` ergänzt (nur Anzeige hier).")
        w("")
        w("| Feld | Roh | Katalog | In `imc_turbines` heute |")
        w("|------|-----|---------|-------------------------|")
        mapped = {
            "EinheitMastrNummer": "ext_unit_key",
            "NameStromerzeugungseinheit": "name",
            "Bruttoleistung": "rated_power_mw (kW→MW)",
            "Inbetriebnahmedatum": "commissioning_date",
            "EinheitBetriebsstatus": "status_label (via Katalog)",
            "Breitengrad": "location",
            "Laengengrad": "location",
        }
        for key in sorted(av01.keys()):
            raw = av01.get(key)
            label = catalog.get(str(raw), "")
            today = mapped.get(key, "—")
            w(f"| `{key}` | {md_escape(raw)} | {md_escape(label) if label else '—'} | {today} |")
        w("")
        extra = sorted(k for k in av01 if k not in mapped)
        w(
            "Zusätzliche gefüllte Keys (nicht im Transform): "
            + ", ".join(f"`{k}`" for k in extra)
            + "."
        )
        w("")
        # manufacturer/type across AV units
        def cat(p, field):
            raw = p.get(field)
            return catalog.get(str(raw), str(raw) if raw is not None else "")

        w("### Hersteller / Typ / Nabe / Rotor über die 12 AV-Einheiten")
        w("")
        w("| Name | SEE | Hersteller | Typ | Nabe m | Rotor m | Brutto kW |")
        w("|------|-----|------------|-----|--------|---------|-----------|")
        for p in sorted(av_units, key=lambda x: str(x.get("NameStromerzeugungseinheit") or "")):
            w(
                "| "
                + " | ".join(
                    [
                        md_escape(p.get("NameStromerzeugungseinheit")),
                        md_escape(p.get("EinheitMastrNummer")),
                        md_escape(cat(p, "Hersteller")),
                        md_escape(p.get("Typenbezeichnung")),
                        md_escape(p.get("Nabenhoehe")),
                        md_escape(p.get("Rotordurchmesser")),
                        md_escape(p.get("Bruttoleistung")),
                    ]
                )
                + " |"
            )
        w("")

    w("## Bewertung (Abbruchkriterium Gewichte)")
    w("")
    spec_weights = []
    meas_weights = []
    for rec in av_specs:
        spec_weights.append(
            any(cell(rec.get(k)) for k in ("RotorWeightt", "BladesWeightt", "TowerWeightt", "NacelleWeightt"))
        )
    for rec in av_meas:
        meas_weights.append(
            any(cell(rec.get(k)) for k in ("RotorWeightt", "BladesWeightt", "TowerWeightt", "NacelleWeightt"))
        )
    w(
        f"- 4C Specs Gewichte gefüllt (mind. ein Weight-Feld): "
        f"{sum(spec_weights)}/{len(av_specs)} Typen."
    )
    w(
        f"- VPI Measurements Gewichte gefüllt: "
        f"{sum(meas_weights)}/{len(av_meas)} Typen."
    )
    w("- MaStR: Typ/Hersteller/Nabe/Rotor **prüfen in Tabelle oben** — das ist der Hebel für Einheiten, unabhängig von 4C-Massen.")
    w("")
    w("## Empfehlung Stufe 1–2")
    w("")
    w("(wird nach dem Lauf im Plan ergänzt)")
    w("")

    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {OUT}")
    print(f"DE01 links={len(av_links)} tids={tids}")
    print(f"AV units scanned={len(av_units)} av01={'yes' if av01 else 'no'}")
    if av01:
        print("AV01 keys:", ", ".join(sorted(av01.keys())))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
